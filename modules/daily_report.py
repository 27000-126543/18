import os
from datetime import datetime, timedelta
from core.logger import logger
from core.database import get_db_cursor
from utils.helpers import today_str, now_str, safe_float, safe_int, op_logger
from config.settings import REPORT_CONFIG
from modules.notifications import NotificationManager


class DailyReportGenerator:
    def __init__(self):
        self.output_dir = REPORT_CONFIG['report_output_dir']
        os.makedirs(self.output_dir, exist_ok=True)

    def export_excel(self, stats, filepath):
        """使用openpyxl生成真正的xlsx格式日报"""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError as e:
            logger.error(f"导出Excel失败: openpyxl未安装 - {e}")
            return False

        try:
            wb = openpyxl.Workbook()

            ws1 = wb.active
            ws1.title = '日报汇总'

            header_font = Font(bold=True, color='FFFFFF', size=11)
            header_fill = PatternFill(start_color='0A1628', end_color='0A1628', fill_type='solid')
            title_font = Font(bold=True, size=14, color='0066FF')
            label_font = Font(bold=True)
            thin_border = Border(
                left=Side(style='thin', color='D0D0D0'),
                right=Side(style='thin', color='D0D0D0'),
                top=Side(style='thin', color='D0D0D0'),
                bottom=Side(style='thin', color='D0D0D0')
            )

            ws1['A1'] = '制造业设备预测性维护与备件管理 - 日报'
            ws1['A1'].font = title_font
            ws1.merge_cells('A1:B1')

            ws1['A2'] = f'报告日期: {stats["report_date"]}'
            ws1['A2'].font = Font(italic=True, color='666666')
            ws1.merge_cells('A2:B2')

            summary_data = [
                ['报告日期', stats['report_date']],
                ['设备总数', stats['total_equipments']],
                ['活跃设备', stats['active_equipments']],
                ['高风险设备', stats['high_risk_equipments']],
                ['今日异常数据点', stats['fault_count']],
                ['设备故障率(%)', stats['fault_rate']],
                ['平均修复时间(小时)', stats['avg_repair_time_hours']],
                ['今日维护成本(¥)', stats['total_maintenance_cost']],
                ['待处理工单', stats['pending_work_orders']],
                ['已完成工单', stats['completed_work_orders']],
                ['低库存备件数', stats['low_stock_items']],
            ]
            for i, row in enumerate(summary_data, start=4):
                c1 = ws1.cell(row=i, column=1, value=row[0])
                c1.font = label_font
                c1.fill = PatternFill(start_color='F0F7FF', end_color='F0F7FF', fill_type='solid')
                c1.alignment = Alignment(horizontal='left', vertical='center')
                c1.border = thin_border
                c2 = ws1.cell(row=i, column=2, value=row[1])
                c2.alignment = Alignment(horizontal='right', vertical='center')
                c2.border = thin_border

            ws1.column_dimensions['A'].width = 22
            ws1.column_dimensions['B'].width = 20

            with get_db_cursor() as cursor:
                cursor.execute(f"""
                    SELECT wo.order_no, wo.title, wo.priority, wo.status,
                           wo.maintenance_cost, wo.created_at, e.equipment_code, e.name
                    FROM work_orders wo
                    JOIN equipments e ON wo.equipment_id = e.id
                    WHERE date(wo.created_at) = ?
                    ORDER BY wo.created_at DESC LIMIT 100
                """, (stats['report_date'],))
                work_orders = [dict(r) for r in cursor.fetchall()]

            if work_orders:
                ws2 = wb.create_sheet('今日工单')
                wo_headers = ['工单号', '标题', '设备编号', '设备名称', '优先级', '状态', '维护成本(¥)', '创建时间']
                for col, header in enumerate(wo_headers, start=1):
                    cell = ws2.cell(row=1, column=col, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border

                for r, wo in enumerate(work_orders, start=2):
                    ws2.cell(row=r, column=1, value=wo['order_no']).border = thin_border
                    ws2.cell(row=r, column=2, value=wo['title']).border = thin_border
                    ws2.cell(row=r, column=3, value=wo['equipment_code']).border = thin_border
                    ws2.cell(row=r, column=4, value=wo['name']).border = thin_border
                    ws2.cell(row=r, column=5, value=wo['priority']).border = thin_border
                    ws2.cell(row=r, column=6, value=wo['status']).border = thin_border
                    ws2.cell(row=r, column=7, value=wo['maintenance_cost']).border = thin_border
                    ws2.cell(row=r, column=8, value=wo['created_at']).border = thin_border

                widths = [18, 30, 14, 20, 8, 10, 12, 20]
                for idx, w in enumerate(widths, start=1):
                    ws2.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

            with get_db_cursor() as cursor:
                cursor.execute("""
                    SELECT sp.part_code, sp.name, sp.current_stock, sp.safety_stock, sp.unit_price
                    FROM spare_parts sp
                    WHERE sp.current_stock <= sp.safety_stock
                    ORDER BY CAST(sp.current_stock AS REAL) / NULLIF(sp.safety_stock, 0) ASC LIMIT 100
                """)
                low_stock = [dict(r) for r in cursor.fetchall()]

            if low_stock:
                ws3 = wb.create_sheet('低库存备件')
                ls_headers = ['备件编码', '备件名称', '当前库存', '安全库存', '单价(¥)']
                for col, header in enumerate(ls_headers, start=1):
                    cell = ws3.cell(row=1, column=col, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border

                for r, part in enumerate(low_stock, start=2):
                    ws3.cell(row=r, column=1, value=part['part_code']).border = thin_border
                    ws3.cell(row=r, column=2, value=part['name']).border = thin_border
                    ws3.cell(row=r, column=3, value=part['current_stock']).border = thin_border
                    ws3.cell(row=r, column=4, value=part['safety_stock']).border = thin_border
                    ws3.cell(row=r, column=5, value=part['unit_price']).border = thin_border

                widths = [16, 30, 12, 12, 12]
                for idx, w in enumerate(widths, start=1):
                    ws3.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

            wb.save(filepath)
            logger.info(f"日报Excel已生成: {filepath}")
            return True
        except Exception as e:
            logger.error(f"导出日报Excel失败: {e}")
            return False

    def export_pdf(self, stats, filepath):
        """使用reportlab生成带表格的真正PDF日报（已注册中文字体）"""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib import colors
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
        except ImportError as e:
            logger.error(f"导出PDF失败: reportlab未安装 - {e}")
            return False

        try:
            font_candidates = [
                ('/System/Library/Fonts/PingFang.ttc', 0, 'PingFangSC'),
                ('/System/Library/Fonts/STHeiti Light.ttc', 0, 'STHeiti'),
                ('/System/Library/Fonts/STHeiti Medium.ttc', 0, 'STHeitiMedium'),
                ('/Library/Fonts/SimSun.ttf', 0, 'SimSun'),
                ('/System/Library/Fonts/Hiragino Sans GB.ttc', 0, 'HiraginoSansGB'),
            ]
            registered_name = 'Helvetica'
            for font_path, sub_idx, font_name in font_candidates:
                if os.path.exists(font_path):
                    try:
                        if font_path.lower().endswith('.ttc'):
                            pdfmetrics.registerFont(TTFont(font_name, font_path, subfontIndex=sub_idx))
                        else:
                            pdfmetrics.registerFont(TTFont(font_name, font_path))
                        registered_name = font_name
                        logger.debug(f"PDF中文字体已注册: {font_name} <- {font_path}")
                        break
                    except Exception as fe:
                        logger.debug(f"尝试注册字体 {font_path} 失败: {fe}")
                        continue

            doc = SimpleDocTemplate(filepath, pagesize=A4)
            styles = getSampleStyleSheet()
            elements = []

            title_style = ParagraphStyle('DailyTitle', parent=styles['Heading1'],
                                         fontName=registered_name,
                                         fontSize=18, textColor=colors.HexColor('#0066FF'),
                                         alignment=1, spaceAfter=6)
            subtitle_style = ParagraphStyle('DailySubtitle', parent=styles['Normal'],
                                            fontName=registered_name,
                                            fontSize=10, textColor=colors.gray,
                                            alignment=1, spaceAfter=20)
            normal_style = ParagraphStyle('DailyNormal', parent=styles['Normal'],
                                          fontName=registered_name, fontSize=10)
            heading_style = ParagraphStyle('DailyH2', parent=styles['Heading2'],
                                           fontName=registered_name, fontSize=13,
                                           textColor=colors.HexColor('#0A1628'),
                                           spaceBefore=8, spaceAfter=10)
            footer_style = ParagraphStyle('Footer', parent=styles['Normal'],
                                          fontName=registered_name,
                                          fontSize=8, textColor=colors.gray, alignment=1)

            elements.append(Paragraph("制造业设备预测性维护与备件管理日报", title_style))
            elements.append(Paragraph(f"报告日期: {stats['report_date']}  |  生成时间: {now_str()}", subtitle_style))

            header_color = colors.HexColor('#0A1628')
            alt_color = colors.HexColor('#F0F7FF')

            summary_data = [
                ['指标', '数值'],
                ['设备总数', str(stats['total_equipments']) + ' 台'],
                ['活跃设备', str(stats['active_equipments']) + ' 台'],
                ['高风险设备', str(stats['high_risk_equipments']) + ' 台'],
                ['今日异常数据点', str(stats['fault_count']) + ' 个'],
                ['设备故障率', f"{stats['fault_rate']:.2f}%"],
                ['平均修复时间', f"{stats['avg_repair_time_hours']:.2f} 小时"],
                ['今日维护成本', f"¥{stats['total_maintenance_cost']:,.2f}"],
                ['待处理工单', str(stats['pending_work_orders']) + ' 个'],
                ['已完成工单', str(stats['completed_work_orders']) + ' 个'],
                ['低库存备件', str(stats['low_stock_items']) + ' 个'],
            ]
            summary_table = Table(summary_data, colWidths=[220, 200])
            summary_style = [
                ('BACKGROUND', (0, 0), (-1, 0), header_color),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
                ('FONTNAME', (0, 0), (-1, -1), registered_name),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                ('TOPPADDING', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D0D0D0')),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
            ]
            for i in range(1, len(summary_data)):
                if i % 2 == 0:
                    summary_style.append(('BACKGROUND', (0, i), (-1, i), alt_color))
            summary_table.setStyle(TableStyle(summary_style))
            elements.append(summary_table)
            elements.append(Spacer(1, 24))

            with get_db_cursor() as cursor:
                cursor.execute(f"""
                    SELECT wo.order_no, wo.title, wo.priority, wo.status
                    FROM work_orders wo
                    WHERE date(wo.created_at) = ?
                    ORDER BY wo.created_at DESC LIMIT 15
                """, (stats['report_date'],))
                work_orders = [dict(r) for r in cursor.fetchall()]

            if work_orders:
                elements.append(Paragraph("今日工单 (Top 15)", heading_style))
                wo_header = ['工单号', '标题', '优先级', '状态']
                wo_data = [wo_header]
                for wo in work_orders:
                    wo_data.append([
                        wo['order_no'],
                        (wo['title'] or '')[:25],
                        wo['priority'],
                        wo['status']
                    ])
                wo_table = Table(wo_data, colWidths=[90, 200, 60, 70])
                wo_style = [
                    ('BACKGROUND', (0, 0), (-1, 0), header_color),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTNAME', (0, 0), (-1, -1), registered_name),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D0D0D0')),
                ]
                for i in range(1, len(wo_data)):
                    if i % 2 == 0:
                        wo_style.append(('BACKGROUND', (0, i), (-1, i), alt_color))
                wo_table.setStyle(TableStyle(wo_style))
                elements.append(wo_table)
                elements.append(Spacer(1, 20))

            with get_db_cursor() as cursor:
                cursor.execute("""
                    SELECT sp.part_code, sp.name, sp.current_stock, sp.safety_stock
                    FROM spare_parts sp
                    WHERE sp.current_stock <= sp.safety_stock
                    ORDER BY CAST(sp.current_stock AS REAL) / NULLIF(sp.safety_stock, 0) ASC LIMIT 15
                """)
                low_stock = [dict(r) for r in cursor.fetchall()]

            if low_stock:
                elements.append(Paragraph("低库存备件预警 (Top 15)", heading_style))
                ls_header = ['备件编码', '备件名称', '当前库存', '安全库存']
                ls_data = [ls_header]
                for part in low_stock:
                    ls_data.append([
                        part['part_code'],
                        (part['name'] or '')[:20],
                        str(part['current_stock']),
                        str(part['safety_stock'])
                    ])
                ls_table = Table(ls_data, colWidths=[100, 180, 70, 70])
                ls_style = [
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D97706')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTNAME', (0, 0), (-1, -1), registered_name),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D0D0D0')),
                ]
                for i in range(1, len(ls_data)):
                    if i % 2 == 0:
                        ls_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#FFFBEB')))
                ls_table.setStyle(TableStyle(ls_style))
                elements.append(ls_table)

            elements.append(Spacer(1, 30))
            elements.append(Paragraph(f"本报告由制造业预测性维护系统自动生成 · {now_str()}", footer_style))

            doc.build(elements)
            logger.info(f"日报PDF已生成: {filepath}")
            return True
        except Exception as e:
            logger.error(f"导出日报PDF失败: {e}")
            return False

    def _collect_daily_stats(self, report_date=None):
        if not report_date:
            report_date = today_str()

        stats = {
            'report_date': report_date,
            'total_equipments': 0,
            'active_equipments': 0,
            'fault_count': 0,
            'fault_rate': 0.0,
            'avg_repair_time_hours': 0.0,
            'total_maintenance_cost': 0.0,
            'high_risk_equipments': 0,
            'pending_work_orders': 0,
            'completed_work_orders': 0,
            'low_stock_items': 0,
        }

        with get_db_cursor() as cursor:
            cursor.execute("SELECT COUNT(*) as cnt FROM equipments")
            stats['total_equipments'] = safe_int(cursor.fetchone()['cnt'])

            cursor.execute("SELECT COUNT(*) as cnt FROM equipments WHERE status = 'active'")
            stats['active_equipments'] = safe_int(cursor.fetchone()['cnt'])

            cursor.execute(f"""
                SELECT COUNT(*) as cnt FROM sensor_data
                WHERE is_anomaly = 1
                AND date(collection_time) = ?
            """, (report_date,))
            stats['fault_count'] = safe_int(cursor.fetchone()['cnt'])

            if stats['active_equipments'] > 0:
                stats['fault_rate'] = round(stats['fault_count'] / stats['active_equipments'] * 100, 2)

            cursor.execute(f"""
                SELECT 
                    AVG((julianday(COALESCE(completed_at, updated_at)) - julianday(created_at)) * 24) as avg_hours,
                    COALESCE(SUM(maintenance_cost), 0) as total_cost,
                    SUM(CASE WHEN status IN ('pending', 'in_progress') THEN 1 ELSE 0 END) as pending_cnt,
                    SUM(CASE WHEN status IN ('completed', 'verified') THEN 1 ELSE 0 END) as completed_cnt
                FROM work_orders
                WHERE date(created_at) = ?
            """, (report_date,))
            row = cursor.fetchone()
            stats['avg_repair_time_hours'] = round(safe_float(row['avg_hours']), 2)
            stats['total_maintenance_cost'] = round(safe_float(row['total_cost']), 2)
            stats['pending_work_orders'] = safe_int(row['pending_cnt'])
            stats['completed_work_orders'] = safe_int(row['completed_cnt'])

            cursor.execute("""
                SELECT COUNT(*) as cnt FROM spare_parts WHERE current_stock <= safety_stock
            """)
            stats['low_stock_items'] = safe_int(cursor.fetchone()['cnt'])

            cursor.execute(f"""
                SELECT COUNT(DISTINCT equipment_id) as cnt FROM equipment_health_records
                WHERE health_status IN ('high_risk', 'medium_risk')
                AND record_date = ?
            """, (report_date,))
            stats['high_risk_equipments'] = safe_int(cursor.fetchone()['cnt'])

        return stats

    def _generate_text_report(self, stats):
        lines = []
        lines.append("=" * 60)
        lines.append(f"  制造业设备预测性维护与备件管理日报")
        lines.append(f"  报告日期: {stats['report_date']}")
        lines.append("=" * 60)
        lines.append("")
        lines.append("【设备概况】")
        lines.append(f"  设备总数: {stats['total_equipments']} 台")
        lines.append(f"  活跃设备: {stats['active_equipments']} 台")
        lines.append(f"  高风险设备: {stats['high_risk_equipments']} 台")
        lines.append("")
        lines.append("【故障统计】")
        lines.append(f"  今日异常数据点: {stats['fault_count']} 个")
        lines.append(f"  设备故障率: {stats['fault_rate']:.2f}%")
        lines.append(f"  平均修复时间: {stats['avg_repair_time_hours']:.2f} 小时")
        lines.append("")
        lines.append("【工单状态】")
        lines.append(f"  待处理工单: {stats['pending_work_orders']} 个")
        lines.append(f"  已完成工单: {stats['completed_work_orders']} 个")
        lines.append(f"  今日维护成本: ¥{stats['total_maintenance_cost']:.2f}")
        lines.append("")
        lines.append("【库存状态】")
        lines.append(f"  低库存备件: {stats['low_stock_items']} 个")
        lines.append("")
        lines.append("=" * 60)
        lines.append(f"  生成时间: {now_str()}")
        lines.append("=" * 60)
        return '\n'.join(lines)

    def generate_daily_report(self, report_date=None):
        if not report_date:
            report_date = today_str()

        logger.info(f"开始生成 {report_date} 日报...")

        stats = self._collect_daily_stats(report_date)

        base_name = f'daily_report_{report_date}'
        pdf_path = os.path.join(self.output_dir, f'{base_name}.pdf')
        excel_path = os.path.join(self.output_dir, f'{base_name}.xlsx')

        self.export_excel(stats, excel_path)
        self.export_pdf(stats, pdf_path)

        try:
            with get_db_cursor(commit=True) as cursor:
                cursor.execute("""
                    INSERT OR REPLACE INTO daily_reports
                    (report_date, total_equipments, active_equipments, fault_count, fault_rate,
                     avg_repair_time_hours, total_maintenance_cost, high_risk_equipments,
                     pending_work_orders, completed_work_orders, low_stock_items,
                     report_path_pdf, report_path_excel)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    report_date, stats['total_equipments'], stats['active_equipments'],
                    stats['fault_count'], stats['fault_rate'], stats['avg_repair_time_hours'],
                    stats['total_maintenance_cost'], stats['high_risk_equipments'],
                    stats['pending_work_orders'], stats['completed_work_orders'],
                    stats['low_stock_items'], pdf_path, excel_path
                ))
        except Exception as e:
            logger.error(f"保存日报记录失败: {e}")

        NotificationManager.notify_daily_report_ready(report_date, pdf_path)

        op_logger.log(
            'generate_daily_report', 'report',
            f'生成 {report_date} 日报, 故障率 {stats["fault_rate"]:.2f}%, 维护成本 ¥{stats["total_maintenance_cost"]:.2f}'
        )

        logger.info(f"日报已生成: PDF={pdf_path}, Excel={excel_path}")
        return stats, pdf_path, excel_path

    def get_recent_reports(self, days=7):
        with get_db_cursor() as cursor:
            cursor.execute(f"""
                SELECT * FROM daily_reports
                WHERE report_date >= date('now', '-{days} days')
                ORDER BY report_date DESC
            """)
            return [dict(r) for r in cursor.fetchall()]
